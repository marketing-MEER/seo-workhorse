"""Shared Excel / CSV writer used by every module."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_report(sheets: dict, out_path) -> Path:
    """Write {sheet_name: DataFrame} to .xlsx (formatted) or .csv (first sheet only)."""
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    if out.suffix.lower() == ".csv":
        next(iter(sheets.values())).to_csv(out, index=False)
        return out

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = name[:31]
            if df.empty:
                df = pd.DataFrame({"result": ["No issues found"]})
            df.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
            ws.auto_filter.ref = ws.dimensions
    return out


def find_column(df: pd.DataFrame, candidates) -> str | None:
    """Return the first column whose lower-cased name matches one of the candidates."""
    lookup = {c.lower().strip(): c for c in df.columns}
    for cand in candidates:
        if cand in lookup:
            return lookup[cand]
    return None


def to_number(series: pd.Series) -> pd.Series:
    """Turn '1,234', '12.5%' or '' into floats."""
    cleaned = series.astype(str).str.replace(",", "", regex=False).str.replace("%", "", regex=False).str.strip()
    return pd.to_numeric(cleaned, errors="coerce").fillna(0)
